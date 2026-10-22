#!/usr/bin/env python3
"""Convierte auditoria.json (auditar_planillas.py) en listados legibles. SOLO LECTURA de las planillas.

Genera, en la raiz del proyecto:
  PARA_AREAS.xlsx / PARA_AREAS.md   -> lo que cada area tiene que completar o confirmar (archivo, hoja, celda, concepto, que falta).
  NO_CUBIERTOS.xlsx                  -> bienes escritos en el F7 que la hoja Carga no incluye (no se habrian cargado).
Uso: python -I armar_listados.py auditoria.json
"""
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AREAS = {"00": "00 HCD", "03": "03 Economia y Finanzas", "05": "05 Obras y Serv. Publicos", "07": "07 Educacion y Des. de la Comunidad",
         "11": "11 Planificacion Estrategica", "13": "13 Mujeres, GyD", "16": "16 Jefatura de Gabinete", "17": "17 Seguridad Ciudadana",
         "25": "25 Desarrollo Productivo"}
SEV = {"falta": "FALTA DATO", "revisar": "CONFIRMAR", "um": "UNIDAD DE MEDIDA", "um2": "CONFIRMAR UNIDAD DE MEDIDA"}


def escribir(ws, cab, filas):
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append(f)
    for i, _ in enumerate(cab, 1):
        ancho = max(len(str(ws.cell(r, i).value or "")) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, ancho + 2), 70)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def guardar(wb, ruta):
    """Si el archivo esta abierto en Excel (bloqueado), guarda con la hora en el nombre en vez de fallar."""
    try:
        wb.save(ruta)
        return ruta
    except PermissionError:
        alt = ruta.with_name(f"{ruta.stem}_{__import__('datetime').datetime.now():%H%M%S}{ruta.suffix}")
        wb.save(alt)
        print(f"{ruta.name} esta abierto: guardado como {alt.name}", file=sys.stderr)
        return alt


def main():
    datos = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    raiz = Path(__file__).parent
    filas, md = [], defaultdict(list)
    for d in datos:
        if d.get("error"):
            continue
        area = AREAS.get(d["subjurisdiccion"], d["subjurisdiccion"])
        for h in d["hallazgos"]:
            fila = [area, d["archivo"], h["hoja"], h["celda"], h.get("concepto") or "", SEV.get(h["sev"], h["sev"]), h["problema"], h["falta"]]
            filas.append(fila)
            md[area].append(fila)
    wb = Workbook()
    escribir(wb.active, ["Area", "Archivo", "Hoja", "Celda", "Concepto", "Tipo", "Problema", "Que falta / que hay que hacer"], filas)
    wb.active.title = "Para las areas"
    guardar(wb, raiz / "PARA_AREAS.xlsx")

    lineas = ["# Pendientes para las areas (planillas F7, presupuesto 2027)", "",
              "Cada fila indica archivo, hoja y celda exacta de la planilla de cada area. Generado con auditar_planillas.py.", ""]
    for area, fs in md.items():
        lineas += [f"## {area}", "", "| Archivo | Hoja | Celda | Concepto | Tipo | Que falta |", "|---|---|---|---|---|---|"]
        for a in fs:
            lineas.append(f"| {a[1]} | {a[2]} | {a[3]} | {a[4]} | {a[5]} | {a[6]}: {a[7]} |")
        lineas.append("")
    (raiz / "PARA_AREAS.md").write_text("\n".join(lineas), encoding="utf-8")

    nc = []
    for d in datos:
        if d.get("error"):
            continue
        for x in d.get("no_cubiertos", []):
            nc.append([AREAS.get(d["subjurisdiccion"], d["subjurisdiccion"]), d["archivo"], x["hoja"], x["fila"], x["denominacion"], x["ff"],
                       x["codigo"], x["um"], x["cantidad"], x["precio"], x["valor"]])
    wb2 = Workbook()
    escribir(wb2.active, ["Area", "Archivo", "Hoja", "Fila", "Denominacion", "FF", "Codigo", "UM", "Cantidad", "Precio", "Costo total"], nc)
    wb2.active.title = "No cubiertos por Carga"
    guardar(wb2, raiz / "NO_CUBIERTOS.xlsx")
    print(f"PARA_AREAS: {len(filas)} filas | NO_CUBIERTOS: {len(nc)} bienes", file=sys.stderr)


if __name__ == "__main__":
    main()
