#!/usr/bin/env python3
"""Audita (SOLO LECTURA) las planillas F7 de formularios/ y arma auditoria.json.

Para cada planilla vigente (la carpeta mas nueva de cada subjurisdiccion) cruza:
  - hojas visibles "F7 comun" y "F7 especial" (lo que escribieron las areas),
  - hoja "Carga" (lo que tipea el .ahk) y
  - "Listado de bienes" (codigo, unidad de medida y precio oficiales).
Reporta, con la celda exacta: datos faltantes, codigos malformados, unidad de medida dudosa, filas #REF! de Carga
(apuntan a filas borradas del F7: no son bienes perdidos) y diferencias de cantidad de bienes.
No modifica ningun archivo.

Uso: python -I auditar_planillas.py [-o auditoria.json]
"""
import json
import re
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
from openpyxl import load_workbook

RAIZ = Path(__file__).parent / "formularios"
COD = re.compile(r"^\d\.\d\.\d\.\d{5}$")
REF = re.compile(r"'(F7 [^']*)'!\$?([A-Z]+)\$?(\d+)")


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()


def err(v):
    return isinstance(v, str) and v.startswith("#")


def vacio(v):
    return v is None or (isinstance(v, str) and not v.strip()) or v == 0


def carpetas_vigentes():
    por_sub = {}
    for d in RAIZ.iterdir():
        m = re.match(r"^(\d\d) - (?:.*-(\d{8}T\d{6}Z)-)?", d.name)   # con o sin sufijo de fecha de la descarga
        if d.is_dir() and m and not d.name.startswith("_"):
            marca = m.group(2) or ""
            if m.group(1) not in por_sub or marca > por_sub[m.group(1)][0]:
                por_sub[m.group(1)] = (marca, d)
    return {k: v[1] for k, v in sorted(por_sub.items())}


def listado(wb):
    por_cod, por_nom = {}, {}
    ws = wb["Listado de bienes"]
    for fila in ws.iter_rows(min_row=2, values_only=True):
        nom, cod, _, um, umn, _, _, precio = (list(fila) + [None] * 8)[:8]
        if cod and isinstance(cod, str) and COD.match(cod.strip()):
            d = {"den": nom, "cod": cod.strip(), "um": um, "umn": umn, "precio": precio}
            por_cod[cod.strip()] = d
            por_nom.setdefault(norm(nom), d)
    return por_cod, por_nom


def hojas_f7(wb):
    com = next(w for w in wb.worksheets if w.title.startswith("F7 com"))
    esp = next(w for w in wb.worksheets if w.title.startswith("F7 esp"))
    return com, esp


def items_visibles(ws, tipo, por_nom, por_cod, hallazgos, archivo):
    n = 0
    for r in range(11, ws.max_row + 1):
        ff, cod, den, um, qty, precio = (ws.cell(r, c).value for c in range(1, 7))
        if norm(den) in ("SUBTOTAL", "TOTAL"):
            continue
        if vacio(den) and vacio(qty) and vacio(cod):
            continue
        if err(den) or err(qty) or err(ff):
            hallazgos.append(dict(archivo=archivo, hoja=ws.title, celda=f"A{r}:G{r}", fila=r, concepto=den,
                                  problema="celda con error en el F7", falta="corregir la fila", sev="falta"))
            continue
        if vacio(qty) and vacio(den):
            continue
        n += 1
        base = dict(archivo=archivo, hoja=ws.title, fila=r, concepto=den)
        if vacio(ff):
            hallazgos.append({**base, "celda": f"A{r}", "problema": "falta fuente de financiamiento",
                              "falta": "indicar la fuente (ej. 110)", "sev": "falta"})
        if vacio(qty) or not isinstance(qty, (int, float)) or qty != int(qty) or qty <= 0:
            hallazgos.append({**base, "celda": f"E{r}", "problema": f"cantidad invalida ({qty!r})",
                              "falta": "cantidad entera mayor a 0", "sev": "falta"})
        if tipo == "comun":
            d_lst = por_nom.get(norm(den))
            if d_lst is not None and (not isinstance(d_lst["precio"], (int, float)) or d_lst["precio"] <= 0):
                hallazgos.append({**base, "celda": f"C{r}", "problema": "el bien esta en el 'Listado de bienes' pero SIN precio (figura en rojo): se carga con precio 0",
                                  "falta": "cargarlo en la hoja F7 especial con codigo, unidad de medida y precio", "sev": "falta"})
            if vacio(den) or norm(den) not in por_nom:
                hallazgos.append({**base, "celda": f"C{r}", "problema": "la denominacion no coincide con el 'Listado de bienes' (codigo 0)",
                                  "falta": "copiar la denominacion exacta del Listado, o cargar el bien en la hoja F7 especial con codigo, unidad y precio",
                                  "sev": "falta"})
        else:
            if vacio(cod) or err(cod) or not COD.match(str(cod).strip()):
                hallazgos.append({**base, "celda": f"B{r}", "problema": f"codigo faltante o malformado ({cod!r})",
                                  "falta": "codigo con formato 0.0.0.00000", "sev": "falta"})
            if vacio(precio) or not isinstance(precio, (int, float)) or precio <= 0:
                hallazgos.append({**base, "celda": f"F{r}", "problema": f"precio faltante o invalido ({precio!r})",
                                  "falta": "precio unitario", "sev": "falta"})
            elif isinstance(precio, float) and abs(precio - round(precio, 2)) > 1e-9:
                hallazgos.append({**base, "celda": f"F{r}", "problema": f"precio con mas de 2 decimales ({precio!r})",
                                  "falta": "confirmar el precio", "sev": "revisar"})
            elif isinstance(precio, float) and precio != int(precio) and round(precio * 1000) == precio * 1000 and precio < 1000:
                hallazgos.append({**base, "celda": f"F{r}", "problema": f"precio {precio!r} parece un punto de miles mal escrito",
                                  "falta": "confirmar si es {:,.0f}".format(precio * 1000), "sev": "revisar"})
            if vacio(um):
                hallazgos.append({**base, "celda": f"D{r}", "problema": "falta la unidad de medida",
                                  "falta": "unidad de medida (inciso 3 = SERVICIO/176, inciso 4 = CADA UNO/3)", "sev": "falta"})
            cs = str(cod).strip() if cod else ""
            if COD.match(cs) and cs not in por_cod:
                inc = cs[0]
                if inc == "3" and norm(um) not in ("SERVICIO",):
                    hallazgos.append({**base, "celda": f"D{r}", "problema": f"inciso 3 con unidad '{um}' (debe ser SERVICIO, 176)",
                                      "falta": "UM 176", "sev": "um"})
                elif inc == "4" and norm(um) not in ("CADA UNO",):
                    hallazgos.append({**base, "celda": f"D{r}", "problema": f"inciso 4 con unidad '{um}' (debe ser CADA UNO, 3)",
                                      "falta": "UM 3", "sev": "um"})
                elif inc == "2":
                    hallazgos.append({**base, "celda": f"D{r}", "problema": f"inciso 2 fuera del listado: unidad '{um}' no verificable",
                                      "falta": "que el area/presupuesto confirme la unidad de medida", "sev": "um2"})
    return n


def auditar(path, codsub):
    wb = load_workbook(path, data_only=True)
    wbf = load_workbook(path, data_only=False)
    nombre = path.name
    hall = []
    por_cod, por_nom = listado(wb)
    com, esp = hojas_f7(wb)
    n_com = items_visibles(com, "comun", por_nom, por_cod, hall, nombre)
    n_esp = items_visibles(esp, "especial", por_nom, por_cod, hall, nombre)
    cv, cf = wb["Carga"], wbf["Carga"]
    validas = errores = 0
    filas_ref = []
    for r in range(7, cv.max_row + 1):
        if str(cv.cell(r, 1).value or "").strip().lower() != "send":
            continue
        vals = [cv.cell(r, c).value for c in (2, 4, 9, 11, 13)]
        forms = [cf.cell(r, c).value for c in (2, 4, 9, 11, 13)]
        if any(err(v) for v in vals):
            errores += 1
            origen = next((REF.search(str(f)) for f in forms if f and REF.search(str(f))), None)
            filas_ref.append(dict(fila_carga=r, origen=(f"{origen.group(1)}!{origen.group(2)}{origen.group(3)}" if origen else "formula con #REF!"),
                                  valores=[str(v) for v in vals]))
            continue
        ff, cod, um, qty, precio = vals
        if all(not vacio(x) for x in (ff, cod, qty, precio)) and COD.match(str(cod).strip()):
            validas += 1
    # Bienes escritos en el F7 que NINGUNA fila de Carga referencia (p. ej. filas insertadas por el area): no se cargarian.
    refs = set()
    for r in range(7, cf.max_row + 1):
        for c in (2, 4, 11, 13):
            m = REF.search(str(cf.cell(r, c).value or ""))
            if m:
                refs.add((m.group(1), int(m.group(3))))
    no_cub = []
    for ws, tipo in ((com, "comun"), (esp, "especial")):
        for r in range(11, ws.max_row + 1):
            ff, cod, den, um, qty, precio = (ws.cell(r, c).value for c in range(1, 7))
            if vacio(den) or norm(den) in ("SUBTOTAL", "TOTAL") or (ws.title, r) in refs:
                continue
            if tipo == "comun":
                d = por_nom.get(norm(den))
                cod, um, precio = (d["cod"], d["um"], d["precio"]) if d else (None, None, None)
            no_cub.append(dict(hoja=ws.title, fila=r, ff=ff, codigo=cod, um=um, cantidad=qty, precio=precio, denominacion=den,
                               valor=(qty * precio) if isinstance(qty, (int, float)) and isinstance(precio, (int, float)) else None))
    return dict(subjurisdiccion=codsub, archivo=nombre, ruta=str(path), items_comun=n_com, items_especial=n_esp,
                no_cubiertos=no_cub,
                items_f7=n_com + n_esp, filas_carga_validas=validas, filas_carga_ref=errores,
                diferencia=(n_com + n_esp) - validas, filas_ref=filas_ref, hallazgos=hall)


def main():
    salida = Path(sys.argv[sys.argv.index("-o") + 1]) if "-o" in sys.argv else Path("auditoria.json")
    res = []
    for sub, carpeta in carpetas_vigentes().items():
        for p in sorted(carpeta.rglob("*.xlsx")):
            if p.name.startswith("~$"):
                continue
            try:
                r = auditar(p, sub)
            except Exception as e:  # una planilla rota no frena la auditoria
                r = dict(subjurisdiccion=sub, archivo=p.name, ruta=str(p), error=f"{type(e).__name__}: {e}")
            res.append(r)
            print(sub, p.name, "->", r.get("error") or f"F7={r['items_f7']} carga_ok={r['filas_carga_validas']} ref={r['filas_carga_ref']} hallazgos={len(r['hallazgos'])}",
                  file=sys.stderr, flush=True)
    salida.write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(res)} planillas -> {salida}", file=sys.stderr)


if __name__ == "__main__":
    main()
