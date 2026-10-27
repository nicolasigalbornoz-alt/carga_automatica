#!/usr/bin/env python3
"""Concilia (SOLO LECTURA) lo que dicen los F7 de la carpeta valida contra lo que se tipeo en RAFAM (cargas/**/*.ahk).

Por cada bien del F7 (hojas "F7 comun" y "F7 especial") busca si algun lote .ahk de su categoria tipeo ese codigo y fuente, y con que
cantidad y precio. Los que no se tipearon se explican con la evidencia disponible (PENDIENTES_REVISION.md, decimales, unidad desconocida).
Nota: lo que ya existia en RAFAM antes de la carga y se salteo por duplicado tampoco aparece en los lotes: figura como "no tipeado".
Genera REVISION_F7_vs_CARGADO.xlsx (+ .md). Uso: python -I conciliar_f7.py
"""
import re
import sys
import warnings
from collections import defaultdict
from pathlib import Path

warnings.filterwarnings("ignore")
sys.path.insert(0, str(Path(__file__).parent))
import auditar_planillas as A
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).parent
SEND = re.compile(r"^Send (\d+)\{Enter\}(\d\.\d\.\d\.\d{5})\{Enter\}\{Delete\}\{Delete\}\{Delete\}(\d+)\{Enter\}([\d.]+)\{Enter\}([\d.]+)")
CATALOGO = {"CADA UNO", "C/U", "UNIDAD", "KILOGRAMO", "LITROS", "METRO", "METRO CUADRADO", "METRO CUBICO", "JUEGO", "TONELADA", "TON", "MES", "PLANCHA",
            "PAR", "FRASCO", "LATA", "SERVICIO", "CAJA", "PACK", "RESMA", "ROLLO", "RL", "DOCENA", "PAQUETE", "BOBINA", "BOLSA", "TAMBOR", "POTE",
            "BALDE", "BIDON X 5 LITROS", "BIDON X 20 LTRS"}


def tipeado():
    t = defaultdict(list)    # (sub, cat) -> [(ff, cod, um, qty, precio, archivo)]
    for p in sorted((RAIZ / "cargas").rglob("*.ahk")):
        m = re.match(r"^(\d\d)_(\d\d\.\d\d)", p.stem)
        clave = (m.group(1), m.group(2)) if m else None
        if not clave and p.stem.startswith("redondeo_01.07"):
            clave = ("03", "01.07")          # filas de 01.07 cargadas el 23/10 con el precio redondeado a 2 decimales
        if not clave:
            continue
        for linea in p.read_text(encoding="utf-8-sig", errors="ignore").splitlines():
            s = SEND.match(linea.strip())
            if s:
                ff, cod, um, qty, pr = s.groups()
                t[clave].append((ff, cod, int(um), float(qty), float(pr), p.name))
    # 01.13 de Economia: se cargo con 03 01.13.ahk (el gas oil 2.5.6.00793 no se tipeo: ya existia)
    for linea in (RAIZ / "03 01.13.ahk").read_text(encoding="utf-8-sig", errors="ignore").splitlines():
        s = SEND.match(linea.strip())
        if s and s.group(2) != "2.5.6.00793":
            ff, cod, um, qty, pr = s.groups()
            t[("03", "01.13")].append((ff, cod, int(um), float(qty), float(pr), "03 01.13.ahk"))
    return t


def ff_txt(v):
    try:
        return str(int(float(v)))
    except (TypeError, ValueError):
        return str(v).strip()


def esperados(path):
    wb = load_workbook(path, data_only=True)
    por_cod, por_nom = A.listado(wb)
    com, esp = A.hojas_f7(wb)
    for ws, tipo in ((com, "F7 comun"), (esp, "F7 especial")):
        for r in range(11, ws.max_row + 1):
            ff, cod, den, um, qty, pr = (ws.cell(r, c).value for c in range(1, 7))
            if A.vacio(den) or A.norm(den) in ("SUBTOTAL", "TOTAL"):
                continue
            if tipo == "F7 comun":
                d = por_nom.get(A.norm(den))
                cod, um, pr = (d["cod"], d["um"], d["precio"]) if d else (None, um, None)
            yield dict(hoja=tipo, fila=r, ff=ff, cod=str(cod).strip() if cod else None, den=den, um=um, qty=qty, precio=pr)


def main():
    pend = (RAIZ / "PENDIENTES_REVISION.md").read_text(encoding="utf-8", errors="ignore")
    t = tipeado()
    faltan, difieren, resumen = [], [], []
    for sub, carpeta in A.carpetas_vigentes().items():
        for p in sorted(carpeta.rglob("*.xlsx")):
            if p.name.startswith("~$"):
                continue
            cat = p.stem[:5]
            usados = t.get((sub, cat), [])
            bolsa = defaultdict(list)
            for x in usados:
                bolsa[(x[0], x[1])].append(x)
            n = ok = 0
            for it in esperados(p):
                q, pr = it["qty"], it["precio"]
                sin_precio_legible = (it["hoja"] == "F7 comun" and it["cod"] and A.COD.match(it["cod"]) and not A.vacio(it["ff"])
                                      and isinstance(q, (int, float)) and q > 0 and not (isinstance(pr, (int, float)) and pr > 0))
                if sin_precio_legible and bolsa.get((ff_txt(it["ff"]), it["cod"])):
                    # el precio sale de formulas del Listado sin valores guardados: no se puede comparar, pero el bien SI se tipeo
                    n += 1
                    ok += 1
                    bolsa[(ff_txt(it["ff"]), it["cod"])].pop()
                    continue
                if not (isinstance(q, (int, float)) and q > 0) or not (isinstance(pr, (int, float)) and pr > 0) or not it["cod"] or not A.COD.match(it["cod"]) or A.vacio(it["ff"]):
                    if isinstance(q, (int, float)) and q > 0:
                        faltan.append([sub, p.name, it["hoja"], it["fila"], it["den"], it["cod"], ff_txt(it["ff"]), q, pr, "DATOS INCOMPLETOS",
                                       "falta codigo, precio o fuente (o el bien no esta en el Listado): lo tienen que completar las areas"])
                    continue
                n += 1
                clave = (ff_txt(it["ff"]), it["cod"])
                if bolsa.get(clave):
                    t_ = min(bolsa[clave], key=lambda x: abs(x[3] - q) + abs(x[4] - pr))
                    bolsa[clave].remove(t_)
                    if abs(t_[3] - q) < 0.5 and abs(t_[4] - pr) < 0.51:
                        ok += 1
                    else:
                        difieren.append([sub, p.name, it["hoja"], it["fila"], it["den"], it["cod"], ff_txt(it["ff"]), q, pr, t_[3], t_[4], t_[5]])
                    continue
                motivo, detalle = "SIN EVIDENCIA", "no se tipeo en ningun lote: probable duplicado ya existente en RAFAM (se salteo) o exclusion no anotada"
                um_txt = A.norm(it["um"]) if it["hoja"] == "F7 especial" else ""
                if abs(pr - round(pr, 2)) > 0.005:
                    motivo, detalle = "PRECIO CON +2 DECIMALES", "RAFAM guarda 2 decimales: se aparto para no cargar mal el precio"
                elif um_txt and um_txt not in CATALOGO:
                    motivo, detalle = f"UNIDAD '{um_txt}' SIN CODIGO", "la unidad no esta en el catalogo del Listado: falta el codigo de UM"
                elif any(it["cod"] in l and cat in l for l in pend.splitlines()):
                    linea = next(l for l in pend.splitlines() if it["cod"] in l and cat in l)
                    motivo = "EXPLICADO EN PENDIENTES"
                    detalle = re.sub(r"\s+", " ", linea)[:170]
                faltan.append([sub, p.name, it["hoja"], it["fila"], it["den"], it["cod"], ff_txt(it["ff"]), q, pr, motivo, detalle])
            resumen.append([sub, p.name, n, ok, sum(1 for f in faltan if f[0] == sub and f[1] == p.name), sum(1 for d in difieren if d[0] == sub and d[1] == p.name)])
    wb = Workbook()
    hojas = [("Resumen", ["Sub", "Archivo", "Bienes con datos", "Tipeados igual", "No tipeados / incompletos", "Tipeados con otro valor"], resumen),
             ("No cargados", ["Sub", "Archivo", "Hoja", "Fila", "Denominacion", "Codigo", "FF", "Cantidad", "Precio", "Motivo", "Detalle"], faltan),
             ("Valores distintos", ["Sub", "Archivo", "Hoja", "Fila", "Denominacion", "Codigo", "FF", "Cant. F7", "Precio F7", "Cant. tipeada", "Precio tipeado", "Lote"], difieren)]
    for i, (nom, cab, filas) in enumerate(hojas):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = nom
        ws.append(cab)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        for f in filas:
            ws.append(f)
        for j in range(1, len(cab) + 1):
            ws.column_dimensions[get_column_letter(j)].width = min(60, max(10, max(len(str(ws.cell(r, j).value or "")) for r in range(1, min(ws.max_row, 80) + 1)) + 2))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    destino = RAIZ / "REVISION_F7_vs_CARGADO.xlsx"
    try:
        wb.save(destino)
    except PermissionError:
        destino = RAIZ / f"REVISION_F7_vs_CARGADO_{__import__('datetime').datetime.now():%H%M%S}.xlsx"
        wb.save(destino)
    por = defaultdict(int)
    for f in faltan:
        por[f[9]] += 1
    print(f"-> {destino.name}")
    print("bienes con datos:", sum(r[2] for r in resumen), "| tipeados igual:", sum(r[3] for r in resumen), "| no cargados:", len(faltan), "| valores distintos:", len(difieren))
    for k, v in sorted(por.items(), key=lambda kv: -kv[1]):
        print(f"   {v:4d}  {k}")


if __name__ == "__main__":
    main()
