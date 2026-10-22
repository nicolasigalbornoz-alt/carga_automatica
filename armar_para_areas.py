#!/usr/bin/env python3
"""Arma PARA_AREAS.xlsx (y PARA_AREAS.md) con lo que las areas todavia tienen que completar o confirmar, segun el estado de la
carga del 22/10/2026. Hoja 1: pedidos a las areas (archivo, hoja, celda). Hoja 2: correcciones internas en RAFAM (decide presupuesto).
Uso: python -I armar_para_areas.py"""
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent

AREAS = [
    # area, archivo, hoja, celda, concepto, tipo, problema, que se necesita
    ("03 Economia y Finanzas", "01.07.00.xlsx", "F7 especial", "F29", "MANTENIMIENTO DE INFRAESTRUCTURA INFORMATICA (3.4.6.08119)", "CONFIRMAR PRECIO",
     "El precio unitario tiene mas de 2 decimales (5.600.699,7858) y RAFAM solo admite 2. NO esta cargado.",
     "Indicar el precio unitario con 2 decimales (12 unidades) para poder cargarlo."),
    ("03 Economia y Finanzas", "01.07.00.xlsx", "F7 especial", "F34", "SERVIDOR (4.3.6.01555)", "CONFIRMAR PRECIO",
     "El precio unitario tiene mas de 2 decimales (525.000,0042). NO esta cargado.",
     "Indicar el precio unitario con 2 decimales (12 unidades) para poder cargarlo."),
    ("11 Planificacion Estrategica", "01.25.00.xlsx", "F7 especial", "F29", "SERVICIO ASESORAMIENTO PROFESIONAL (3.4.9.04331)", "CONFIRMAR PRECIO",
     "La planilla dice 12 x $6.500.000, pero en RAFAM esta categoria ya tiene ese codigo (fuente 110) con 12 x $3.359.682. NO se cargo para no duplicar.",
     "Confirmar cual es el precio correcto; si es el de la planilla, presupuesto corrige el existente en RAFAM."),
    ("11 Planificacion Estrategica", "01.25.00.xlsx", "F7 especial", "F22", "MANT. EQUIPO DE COMPUTACION (3.3.3.01657)", "CONFIRMAR PRECIO",
     "El precio unitario tiene decimales de mas (521.785,71428571). Se cargo, y RAFAM lo guarda con 2 decimales.",
     "Confirmar el precio unitario con 2 decimales (14 servicios)."),
    ("16 Jefatura de Gabinete", "01.06.00.xlsx", "F7 comun", "C40", "NEUMATICO P/AUTO (cantidad 10)", "FALTA DATO",
     "El bien no figura en el Listado de bienes (codigo y precio = 0). NO esta cargado.",
     "Indicar codigo del bien (0.0.0.00000), unidad de medida y precio, cargandolo en la hoja F7 especial."),
    ("16 Jefatura de Gabinete", "35.01.00.xlsx", "F7 especial", "B16", "ESMALTE ANTIOXIDANTE BLANCO", "FALTA DATO",
     "Codigo malformado '2.5.5.0057' (4 digitos finales en vez de 5; no esta en el Listado). NO esta cargado.",
     "Indicar el codigo correcto (¿2.5.5.00570 ESMALTE SINTETICO u otro?)."),
    ("16 Jefatura de Gabinete", "48.00.00.xlsx", "F7 especial", "F24:F30", "ESTANTE, PERCHERO, SILLA, MESA DE REUNIONES, VENTILADOR DE PARED, BIBLIOTECA, CANASTO METALICO",
     "CONFIRMAR PRECIO", "Los precios son muy bajos ($80, $40, $100, $350, $100, $200 y $20 por unidad). Se cargaron tal cual; parecen estar expresados en miles.",
     "Confirmar el precio unitario real de cada bien (y revisar tambien los otros precios muy bajos de la planilla)."),
    ("07 Educacion y Des. de la Comunidad", "42.01.00.xlsx", "F7 especial", "B11", "CORTINA BLACK OUT (cant. 7, $400.000)", "FALTA DATO",
     "Sin codigo del bien. NO esta cargado.", "Indicar el codigo del bien (0.0.0.00000)."),
    ("07 Educacion y Des. de la Comunidad", "42.01.00.xlsx", "F7 especial", "D11", "CORTINA BLACK OUT", "FALTA DATO",
     "Sin unidad de medida.", "Indicar la unidad de medida (CADA UNO, METRO, etc.)."),
    ("07 Educacion y Des. de la Comunidad", "42.01.00.xlsx", "F7 especial", "B12", "PANTALLA PROYECTOR (cant. 1, $160.000)", "FALTA DATO",
     "Sin codigo del bien. NO esta cargado.", "Indicar el codigo del bien (0.0.0.00000)."),
    ("07 Educacion y Des. de la Comunidad", "42.01.00.xlsx", "F7 especial", "D12", "PANTALLA PROYECTOR", "FALTA DATO",
     "Sin unidad de medida.", "Indicar la unidad de medida."),
]

INTERNO = [
    ("25 Desarrollo Productivo", "16.02 Promocion del Comercio", "SOBRE C/MEMBRETE (2.3.3.00861)", "Cargado a $23,11 (x15 = $346,65) por un punto de miles mal escrito; la planilla ya trae 23113.",
     "Corregir el precio en RAFAM a $23.113 con Modificar + Guardar (decide presupuesto; hoy solo se autorizo modificar 48.00)."),
    ("16 Jefatura de Gabinete", "35.01 Arte y Cultura", "LENTEJUELAS (2.2.9.06109) y PAPEL PASPARTOUF (2.3.1.04846)", "Cargados con UM 3; el area expreso PAQUETE (210).", "Decidir si se corrige la UM en RAFAM."),
    ("16 Jefatura de Gabinete", "35.01 Arte y Cultura", "TELA ACETATO (2.2.9.07097) y PLASTICO DE BURBUJAS (2.5.8.04490)", "Cargados con UM 3; el area expreso ROLLO (208).", "Decidir si se corrige la UM en RAFAM."),
    ("07 Educacion y Des. de la Comunidad", "39.00 Casa Comunal CIS", "Bienes con fuente 110 y 132", "Hay bienes preexistentes con fuente 132 y los cargados con fuente 110 (resuelto: la fuente correcta es 110).",
     "Si corresponde, decidir si se retiran las filas de fuente 132 (borrar lo decide el usuario)."),
]


def hoja(ws, cab, filas):
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append(list(f))
    for i in range(1, len(cab) + 1):
        ws.column_dimensions[get_column_letter(i)].width = [26, 18, 14, 10, 44, 18, 60, 52][i - 1] if len(cab) == 8 else [26, 28, 52, 60, 60][i - 1]
        for r in range(2, ws.max_row + 1):
            ws.cell(r, i).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    wb = Workbook()
    hoja(wb.active, ["Area", "Archivo", "Hoja", "Celda", "Concepto", "Tipo", "Problema", "Que falta / que hay que hacer"], AREAS)
    wb.active.title = "Para las areas"
    hoja(wb.create_sheet("Interno (RAFAM)"), ["Subjurisdiccion", "Categoria", "Bien", "Situacion", "Que hay que decidir"], INTERNO)
    ruta = RAIZ / "PARA_AREAS.xlsx"
    try:
        wb.save(ruta)
    except PermissionError:
        ruta = ruta.with_name(f"PARA_AREAS_{datetime.now():%H%M%S}.xlsx")
        wb.save(ruta)
        print(f"PARA_AREAS.xlsx esta abierto en Excel: guardado como {ruta.name}", file=sys.stderr)

    md = ["# Pendientes para las areas (presupuesto 2027) - actualizado 22/10/2026", "",
          "Todo lo demas de las planillas ya esta cargado en RAFAM. Falta solo lo siguiente.", ""]
    area_prev = None
    for a in AREAS:
        if a[0] != area_prev:
            md += ["", f"## {a[0]}", "", "| Archivo | Hoja | Celda | Concepto | Tipo | Que falta |", "|---|---|---|---|---|---|"]
            area_prev = a[0]
        md.append(f"| {a[1]} | {a[2]} | {a[3]} | {a[4]} | {a[5]} | {a[6]} {a[7]} |")
    (RAIZ / "PARA_AREAS.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print(f"{len(AREAS)} pedidos a las areas, {len(INTERNO)} decisiones internas -> {ruta.name}")


if __name__ == "__main__":
    main()
