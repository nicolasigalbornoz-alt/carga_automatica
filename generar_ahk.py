#!/usr/bin/env python3
"""Genera el .ahk de carga para RAFAM a partir de la hoja de carga del formulario.

Reemplaza los pasos manuales (copiar la hoja, pegarla en el Notepad, Ctrl+H para
sacar los espacios, borrar filas vacias). Uso:

    python generar_ahk.py "01.13.xlsx" -o "03 01.13.ahk"
    python generar_ahk.py hoja.csv --pausa 100 --ventana "RAFAM"
    python generar_ahk.py "03 01.13.ahk"            # re-limpia un .ahk ya armado

Entrada .xlsx/.xlsm: toma la hoja oculta (la de carga) salvo que se indique --hoja.
Cada fila de la hoja es: Send | celdas... -> se concatenan sin espacios.
Se descartan las filas sin Send y las que no tienen ningun dato (vacias o en 0).
En el .ahk generado: F1 carga, F2 frena (ExitApp), Pause pausa/reanuda.
Compatible con AutoHotkey 1.x.
"""
import argparse
import csv
import re
import sys
from pathlib import Path

TOKEN = re.compile(r"^\{.*\}$")  # {Enter}, {Delete}, {Sleep 50}


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def filas_xlsx(path, hoja):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=False)
    if hoja:
        ws = wb[hoja]
    else:
        ocultas = [w for w in wb.worksheets if w.sheet_state != "visible"]
        ws = ocultas[0] if ocultas else (wb["Carga"] if "Carga" in wb.sheetnames else wb.active)
    print(f"hoja: {ws.title}", file=sys.stderr)
    for row in ws.iter_rows(values_only=True):
        yield [fmt(c) for c in row]


def filas_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t") if sample else csv.excel
        for row in csv.reader(f, dialect):
            yield [c.strip() for c in row]


def filas_ahk(path):
    for linea in Path(path).read_text(encoding="utf-8-sig").splitlines():
        m = re.match(r"\s*Send\s+(.*\S)\s*$", linea)
        if m:
            yield ["Send", m.group(1)]


def recortar(filas):
    """Corta cada fila en la ultima columna con una tecla {..}: asi no se cuelan
    columnas auxiliares de la hoja (descripciones, indices) a la derecha del bloque."""
    filas = list(filas)
    ancho = 0
    for f in filas:
        if f and f[0].lower() == "send":
            for i, c in enumerate(f):
                if TOKEN.match(c):
                    ancho = max(ancho, i + 1)
    if ancho:
        print(f"bloque Send: columnas 1-{ancho}", file=sys.stderr)
        filas = [f[:ancho] for f in filas]
    return filas


def cuerpo(celdas):
    """Devuelve el texto tras 'Send ' o None si la fila no sirve."""
    if not celdas or celdas[0].lower() != "send":
        return None
    resto = [c for c in celdas[1:] if c]
    datos = [c for c in resto if not TOKEN.match(c)]
    if not any(d not in ("0", "0.0") for d in datos):
        return None
    return "".join(resto)


CABECERA = """\
#NoEnv  ; Recommended for performance and compatibility with future AutoHotkey releases.
; #Warn  ; Enable warnings to assist with detecting common errors.
SetWorkingDir %A_ScriptDir%  ; Ensures a consistent starting directory.

; Comenzara a tipear automaticamente al pulsar F1
; F2 frena la carga de inmediato (tambien: click derecho en el icono H > Exit)
; Pause pausa / reanuda
; PAUSA: milisegundos extra entre filas (subir si RAFAM se queda atras)
; VENTANA: si no esta vacio, frena si esa ventana no esta activa (titulo parcial)
PAUSA := {pausa}
VENTANA := "{ventana}"
TOTAL := {total}
FILA := 0
{auto}
F2::
ToolTip, Carga frenada en la fila %FILA% de %TOTAL%
Sleep, 1500
ExitApp
return

Pause::Pause

Chk() {{
    global PAUSA, VENTANA, TOTAL, FILA
    FILA++
    if (VENTANA != "" && !WinActive(VENTANA)) {{
        ; el foco puede perderse un instante (notificacion, etc.): reintenta unos segundos antes de frenar
        WinActivate, %VENTANA%
        WinWaitActive, %VENTANA%,, 3
        if (!WinActive(VENTANA)) {{
            MsgBox, 48, Carga frenada, La ventana "%VENTANA%" no esta activa. Fila %FILA% de %TOTAL%.
            ExitApp
        }}
    }}
    ToolTip, Fila %FILA% de %TOTAL%  (F2 frena)
    Sleep, %PAUSA%
}}

{entrada}
"""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("entrada")
    ap.add_argument("-o", "--salida")
    ap.add_argument("--hoja", help="nombre de la hoja (por defecto la oculta)")
    ap.add_argument("--pausa", type=int, default=0, help="ms extra entre filas")
    ap.add_argument("--ventana", default="", help="titulo (parcial) de la ventana de RAFAM a exigir activa")
    ap.add_argument("--primeras", type=int, default=0, help="solo N filas (desde --desde); para cargar por lotes")
    ap.add_argument("--desde", type=int, default=1, help="primera fila a incluir (1 = la primera), luego de --excluir")
    ap.add_argument("--excluir", default="", help="codigos de bien a saltear, separados por coma (ej. 2.5.6.00793)")
    ap.add_argument("--excluir-filas", default="", dest="excluir_filas",
                    help="numeros de fila (1 = la primera fila con datos, ANTES de --excluir) a saltear, separados por coma")
    ap.add_argument("--um", default="",
                    help="unidad de medida a forzar por codigo, separados por coma (ej. 2.5.9.03225=22,3.9.1.07239=176)")
    ap.add_argument("--auto", type=int, default=None, metavar="SEG",
                    help="empieza sola a los SEG segundos de abrir el .ahk, sin esperar F1 (F2 la cancela)")
    a = ap.parse_args()

    ext = Path(a.entrada).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        filas = filas_xlsx(a.entrada, a.hoja)
    elif ext == ".ahk":
        filas = filas_ahk(a.entrada)
    else:
        filas = filas_csv(a.entrada)

    envios = [b for b in (cuerpo(f) for f in recortar(filas)) if b]
    if not envios:
        sys.exit("no se encontraron filas con datos")
    for par in (x for x in a.um.split(",") if x.strip()):
        cod, um = (s.strip() for s in par.split("="))
        pat = re.compile(r"(\{Enter\}" + re.escape(cod) + r"\{Enter\}(?:\{Delete\})*)\d+(\{Enter\})")
        envios = [pat.sub(lambda m: m.group(1) + um + m.group(2), b) for b in envios]
    nf = {int(x) for x in a.excluir_filas.split(",") if x.strip()}
    if nf:
        envios = [b for i, b in enumerate(envios, 1) if i not in nf]
        print(f"excluidas {len(nf)} filas por numero", file=sys.stderr)
    excluir = [c.strip() for c in a.excluir.split(",") if c.strip()]
    if excluir:
        antes = len(envios)
        envios = [b for b in envios if not any(f"{{Enter}}{c}{{Enter}}" in b for c in excluir)]
        print(f"excluidas {antes - len(envios)} filas ({', '.join(excluir)})", file=sys.stderr)
    envios = envios[max(a.desde, 1) - 1:]
    if a.primeras > 0:
        envios = envios[: a.primeras]
    if not envios:
        sys.exit("no quedan filas para cargar con esas opciones")

    salida = Path(a.salida) if a.salida else Path(a.entrada).with_suffix(".ahk")
    if a.auto is None:
        auto, entrada = "", "F1::"
    else:
        auto = (f"AUTO := {max(a.auto, 0)}\n"
                "ToolTip, Carga automatica en %AUTO% s (F2 la cancela)\n"
                "Sleep, % AUTO * 1000\n"
                "GoSub, Cargar\n"
                "ExitApp\n")
        entrada = "Cargar:"
    lineas = CABECERA.format(pausa=a.pausa, ventana=a.ventana.replace('"', ""), total=len(envios),
                             auto=auto, entrada=entrada).splitlines()
    for b in envios:
        lineas += ["Chk()", f"Send {b}"]
    lineas += ["ToolTip", f'MsgBox, 64, Carga terminada, Se enviaron %TOTAL% filas.', "return"]
    salida.write_bytes(("\r\n".join(lineas) + "\r\n").encode("utf-8-sig"))
    print(f"{len(envios)} filas -> {salida}", file=sys.stderr)


if __name__ == "__main__":
    main()
