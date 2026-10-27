#!/usr/bin/env python3
"""Arma (solo lectura) la lista COMPLETA de bienes a cargar por planilla: filas validas de la hoja Carga + bienes del F7 que Carga
no incluye o trae rotos. Valida, resuelve la UM (inc.3=176, inc.4=3, inc.2 = Listado o col. D del area) y concilia con el total del F7.

  python -I prep_26_10.py <planilla.xlsx> <pref> [--portal 60127470]      escribe cargas/f26/<pref>.json y muestra el resumen
"""
import argparse, json, re, sys, warnings
from pathlib import Path
warnings.filterwarnings("ignore")
sys.path.insert(0, str(Path(__file__).resolve().parent))
import validar_um as v
import auditar_planillas as A
from openpyxl import load_workbook

DIR = Path(__file__).resolve().parent
COD = re.compile(r"^\d\.\d\.\d\.\d{5}$")
v.UM_EXTRA.setdefault("UNIDAD", 3)       # decision del usuario 26/10: "UNIDAD" = 3
v.UM_EXTRA.setdefault("PLIEGO", 3)       # decision del usuario 27/10: "PLIEGO" sin codigo en RAFAM -> UM 3 (anotar en PENDIENTES)
REDONDEAR = "--redondear" in sys.argv   # 27/10: precios con mas de 2 decimales -> se redondean a 2 y se cargan


def catalogo_global():
    """nombre de UM -> codigo, uniendo los 'Listado de bienes' de todas las planillas vigentes (solo los nombres sin conflicto)."""
    cat = {}
    for carp in A.carpetas_vigentes().values():
        for p in carp.rglob("*.xlsx"):
            if p.name.startswith("~$") or p.parent.name == "backend":
                continue
            try:
                ws = load_workbook(p, data_only=True, read_only=True)["Listado de bienes"]
            except Exception:
                continue
            for r in ws.iter_rows(min_row=2, max_col=5, values_only=True):
                if r[3] and r[4] is not None:
                    try:
                        cat.setdefault(A.norm(r[3]), set()).add(int(r[4]))
                    except (TypeError, ValueError):
                        pass
    return {k: next(iter(s)) for k, s in cat.items() if len(s) == 1}


_CAT = {}


def cat_get(nombre):
    if not _CAT:
        cache = DIR / "cargas" / "f26" / "_catalogo_um.json"
        if cache.exists():
            _CAT.update(json.loads(cache.read_text(encoding="utf-8")))
        else:
            _CAT.update(catalogo_global())
            cache.parent.mkdir(parents=True, exist_ok=True)
            cache.write_text(json.dumps(_CAT, ensure_ascii=False, indent=1), encoding="utf-8")
    return _CAT.get(nombre)


def f7_filas(wb):
    """Todos los bienes escritos en el F7 (comun y especial): dicts con ff, cod, um_area, cant, precio, den, hoja, fila."""
    por_cod, por_nom = A.listado(wb)
    com, esp = A.hojas_f7(wb)
    out = []
    for ws, tipo in ((com, "comun"), (esp, "especial")):
        for r in range(11, ws.max_row + 1):
            ff, cod, den, um, qty, precio = (ws.cell(r, c).value for c in range(1, 7))
            if A.vacio(den) or A.norm(den) in ("SUBTOTAL", "TOTAL"):
                continue
            if tipo == "comun":
                d = por_nom.get(A.norm(den))
                cod, um, precio = (d["cod"], d["um"], d["precio"]) if d else (None, None, None)
            out.append(dict(hoja=ws.title, fila=r, ff=ff, cod=cod, um_area=um, cant=qty, precio=precio, den=den, tipo=tipo))
    return out


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def clave(ff, cod, cant, precio):
    f, c, q, p = str(ff).strip(), str(cod).strip(), num(cant), num(precio)
    return (re.sub(r"\.0$", "", f), c, None if q is None else round(q), None if p is None else round(p, 2))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("pref")
    ap.add_argument("--portal", type=float, default=None)
    ap.add_argument("--redondear", action="store_true")
    a = ap.parse_args()

    wb = load_workbook(a.xlsx, data_only=True)
    lst, areas = v.listado(a.xlsx), v.um_areas(a.xlsx)
    cat = None
    f7 = f7_filas(wb)
    total_f7 = sum(f["cant"] * f["precio"] for f in f7
                   if num(f["cant"]) and num(f["precio"]))

    # 1) filas de Carga que sirven
    carga, rotas = [], []
    for i, b in enumerate(v.envios(a.xlsx), 1):
        m = v.ROW.match(b)
        if not m:
            rotas.append((i, b))
            continue
        ff, cod, um, cant, precio = m.groups()
        if "#" in b or not (re.fullmatch(r"\d{3}", ff) and COD.match(cod) and re.fullmatch(r"\d+", cant) and num(precio)):
            rotas.append((i, b))
            continue
        carga.append(dict(origen=f"Carga#{i}", ff=ff, cod=cod, um_carga=um, cant=int(cant), precio=float(precio)))
    # 2) bienes del F7 que ninguna fila buena de Carga cubre (multiconjunto por ff, codigo, cantidad, precio)
    pool = {}
    for c in carga:
        pool.setdefault(clave(c["ff"], c["cod"], c["cant"], c["precio"]), []).append(c)
    extras, sin_cod = [], []
    for f in f7:
        if not (num(f["cant"]) and num(f["precio"])):
            sin_cod.append((f, "cantidad o precio vacios/invalidos"))
            continue
        k = clave(f["ff"], f["cod"], f["cant"], f["precio"])
        if pool.get(k):
            pool[k].pop()
        else:
            extras.append(dict(origen=f"{f['hoja']}!{f['fila']}", ff=re.sub(r"\.0$", "", str(f["ff"]).strip()), cod=str(f["cod"]).strip() if f["cod"] else "",
                               um_carga=None, cant=num(f["cant"]), precio=num(f["precio"]), den=f["den"], um_area=A.norm(f["um_area"]), tipo=f["tipo"]))
    huerfanas = [c for lst_ in pool.values() for c in lst_]      # filas de Carga sin bien en el F7

    # una fila de Carga sin bien igual en el F7 es basura (codigo de una fila con cantidad/precio de otra): NO se carga
    ids = {id(c) for c in huerfanas}
    carga = [c for c in carga if id(c) not in ids]

    # 3) validacion y UM
    ok, pend, vistos = [], [], {}
    for r in carga + extras:
        mot = None
        cod, ff = r["cod"], r["ff"]
        pr = r["precio"]
        if abs(pr - round(pr, 2)) < 1e-6:
            pr = round(pr, 2)
        elif abs(pr - round(pr)) < 1e-6:
            pr = float(round(pr))
        if REDONDEAR and abs(pr - round(pr, 2)) > 1e-9:
            r["precio_orig"] = pr
            pr = round(pr, 2)
        r["precio"] = pr
        cant = r["cant"]
        if not re.fullmatch(r"\d{3}", ff):
            mot = f"FF invalida '{ff}'"
        elif not COD.match(cod):
            mot = f"codigo faltante o malformado '{cod}'"
        elif cant is None or cant != int(cant) or cant <= 0:
            mot = f"cantidad invalida '{cant}'"
        elif abs(pr - round(pr, 2)) > 1e-9:
            mot = f"precio con mas de 2 decimales ({pr})"
        elif re.fullmatch(r"\d+\.\d{3}", repr(pr)):
            mot = f"precio {pr}: probable separador de miles"
        else:
            r["cant"] = int(cant)
            inc = cod[0]
            if inc == "4":
                um = 3
            elif inc == "3":
                um = 176
            elif cod in lst:
                um = lst[cod]
            elif cod in areas and areas[cod][0] is not None:
                um = areas[cod][0]
            elif cod in areas and cat_get(areas[cod][1]) is not None:
                um = cat_get(areas[cod][1])
                r["um_nota"] = f"UM '{areas[cod][1]}' tomada del catalogo de otras planillas"
            elif r.get("um_carga") and r["origen"].startswith("F7 com"):
                um = int(r["um_carga"])
            else:
                um = None
                mot = f"inciso 2 fuera del listado: UM del area '{areas.get(cod, (None, '?'))[1]}' no esta en el catalogo"
            r["um"] = um
        if not mot and (ff, cod) in vistos:
            mot = f"repetido en la planilla (ya en {vistos[(ff, cod)]})"
        if mot:
            pend.append((r, mot))
        else:
            vistos[(ff, cod)] = r["origen"]
            ok.append(r)

    tot_ok = sum(r["cant"] * r["precio"] for r in ok)
    print(f"{a.pref}: F7 {len(f7)} bienes / ${total_f7:,.0f} | Carga buenas {len(carga)} + F7 no cubiertos/rotos {len(extras)}"
          f" | a cargar {len(ok)} / ${tot_ok:,.2f} | pendientes {len(pend)} | Carga rotas {len(rotas)} | Carga sin bien en F7 {len(huerfanas)}")
    if a.portal is not None:
        print(f"   portal ${a.portal:,.0f} | F7 {'=' if abs(total_f7 - a.portal) < 1 else '<>'} portal | cargables+pendientes = ${tot_ok + sum(r['cant'] * r['precio'] for r, _ in pend if isinstance(r['cant'], (int, float))):,.2f}")
    for r, m in pend:
        print("   PEND", r["origen"], r["cod"], r.get("den", ""), r["ff"], r["cant"], r["precio"], "->", m)
    for r in huerfanas:
        print("   HUERFANA (Carga sin F7)", r["origen"], r["cod"], r["cant"], r["precio"])
    for r in ok:
        if r.get("um_nota"):
            print("   NOTA", r["cod"], r["um_nota"])
    forz = [(r["cod"], r["um"]) for r in ok if r.get("um_carga") and r["um"] != int(r["um_carga"])]
    if forz:
        print("   UM forzadas respecto de Carga:", forz[:60], "..." if len(forz) > 60 else "")
    out = DIR / "cargas" / "f26"
    out.mkdir(parents=True, exist_ok=True)
    (out / f"{a.pref}.json").write_text(json.dumps(dict(xlsx=a.xlsx, ok=ok, pend=[dict(r, motivo=m) for r, m in pend], total_f7=total_f7),
                                                    ensure_ascii=False, indent=1, default=str), encoding="utf-8")


if __name__ == "__main__":
    main()
